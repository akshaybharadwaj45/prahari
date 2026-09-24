import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import os
import win32com.client

# Base paths
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(BASE_DIR, 'docs')
ASSETS_DIR = os.path.join(DOCS_DIR, 'template_assets')
PREVIEWS_DIR = os.path.join(DOCS_DIR, 'slide_previews')
os.makedirs(DOCS_DIR, exist_ok=True)
os.makedirs(ASSETS_DIR, exist_ok=True)
os.makedirs(PREVIEWS_DIR, exist_ok=True)

bg_title_slide = os.path.join(ASSETS_DIR, 'bg_title_slide.png')
bg_content_slide = os.path.join(ASSETS_DIR, 'bg_content_slide.png')

# -------------------------------------------------------------
# PRESENTATION INITIALIZATION
# -------------------------------------------------------------
prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

COLOR_BG          = RGBColor(255, 255, 255)
COLOR_BLACK       = RGBColor(0, 0, 0)         # Exact 0x0 from template
COLOR_TITLE_RED   = RGBColor(255, 0, 0)       # Exact 0xff0000 from template
COLOR_BLUE_TITLE  = RGBColor(49, 133, 156)    # Exact 0x31859c from template
COLOR_BORDER      = RGBColor(0, 0, 0)         # Black border for boxes

def apply_slide_background(slide, is_title=False):
    bg_img = bg_title_slide if is_title else bg_content_slide
    if os.path.exists(bg_img):
        slide.shapes.add_picture(bg_img, Inches(0), Inches(0), Inches(13.333), Inches(7.5))

def add_header(slide, title_text):
    apply_slide_background(slide, is_title=False)
    
    # Title Text placed at the exact template header position (x=1.07", y=0.36")
    tb = slide.shapes.add_textbox(Inches(1.07), Inches(0.36), Inches(10.5), Inches(0.55))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.name = 'Arial'
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLOR_TITLE_RED

def render_section(tf, heading, items, first=False, head_size=20, bullet_size=15.5, space_head=12, space_item=3.5):
    p_h = tf.paragraphs[0] if first else tf.add_paragraph()
    p_h.text = heading
    p_h.font.name = 'Arial'
    p_h.font.size = Pt(head_size)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_BLACK
    if not first:
        p_h.space_before = Pt(space_head)
        
    for it in items:
        p = tf.add_paragraph()
        p.text = '•  ' + it
        p.font.name = 'Arial'
        p.font.size = Pt(bullet_size)
        p.font.color.rgb = COLOR_BLACK
        p.space_before = Pt(space_item)

# -------------------------------------------------------------
# SLIDE 1: TITLE SLIDE (Uses authentic original master background)
# -------------------------------------------------------------
s1 = prs.slides.add_slide(prs.slide_layouts[6])
apply_slide_background(s1, is_title=True)

# Project Title (Placed at exact y=3.26" matching template)
tb_proj = s1.shapes.add_textbox(Inches(1.5), Inches(3.26), Inches(10.333), Inches(0.8))
tf_proj = tb_proj.text_frame
tf_proj.word_wrap = True

p_p = tf_proj.paragraphs[0]
p_p.text = 'G-1400 | PRAHARI: Autonomous Spacecraft Conjunction Assessment & High-Recall Collision Triage'
p_p.font.name = 'Arial'
p_p.font.size = Pt(20)
p_p.font.bold = True
p_p.font.color.rgb = COLOR_BLUE_TITLE
p_p.alignment = PP_ALIGN.CENTER

# Date (Placed at exact y=4.12" matching template)
tb_date = s1.shapes.add_textbox(Inches(4.5), Inches(4.12), Inches(4.333), Inches(0.4))
tf_date = tb_date.text_frame
p_date = tf_date.paragraphs[0]
p_date.text = 'Date: 25-09-2026'
p_date.font.name = 'Arial'
p_date.font.size = Pt(20)
p_date.font.bold = True
p_date.font.color.rgb = COLOR_BLUE_TITLE
p_date.alignment = PP_ALIGN.CENTER

# Left Box: Team Members (Exact position: x=2.22", y=4.79", w=5.27", h=2.36")
box_team = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.22), Inches(4.79), Inches(5.27), Inches(2.36))
box_team.fill.solid()
box_team.fill.fore_color.rgb = COLOR_BG
box_team.line.color.rgb = COLOR_BORDER
box_team.line.width = Pt(1.5)
tf_t = box_team.text_frame
tf_t.word_wrap = True

p_th = tf_t.paragraphs[0]
p_th.text = 'Team Members'
p_th.font.name = 'Arial'
p_th.font.size = Pt(18)
p_th.font.bold = True
p_th.font.underline = True
p_th.font.color.rgb = COLOR_BLACK
p_th.alignment = PP_ALIGN.CENTER

members = [
    '1.  Akshay Bharadwaj       -24BD1A052K',
    '2.  M. Lakshmi Srivani      - 24BD1A0533',
    '3.  P. Pragna                     - 24BD1A053B',
    '4.  R. Lahari                      - 24BD1A053N',
    '5.  S. Meghana                 - 24BD1A053R'
]
for m in members:
    p_m = tf_t.add_paragraph()
    p_m.text = m
    p_m.font.name = 'Arial'
    p_m.font.size = Pt(15.5)
    p_m.font.color.rgb = COLOR_BLACK
    p_m.space_before = Pt(2)

# Right Box: Mentors (Exact position: x=7.69", y=4.79", w=4.60", h=2.36")
box_men = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.69), Inches(4.79), Inches(4.60), Inches(2.36))
box_men.fill.solid()
box_men.fill.fore_color.rgb = COLOR_BG
box_men.line.color.rgb = COLOR_BORDER
box_men.line.width = Pt(1.5)
tf_men = box_men.text_frame
tf_men.word_wrap = True

p_mh = tf_men.paragraphs[0]
p_mh.text = 'Mentors'
p_mh.font.name = 'Times New Roman'
p_mh.font.size = Pt(20)
p_mh.font.bold = True
p_mh.font.underline = True
p_mh.font.color.rgb = COLOR_BLACK
p_mh.alignment = PP_ALIGN.CENTER

p_mn1 = tf_men.add_paragraph()
p_mn1.text = 'T Rupa Devi Ma’am'
p_mn1.font.name = 'Times New Roman'
p_mn1.font.size = Pt(20)
p_mn1.font.bold = True
p_mn1.font.color.rgb = COLOR_BLACK
p_mn1.alignment = PP_ALIGN.CENTER
p_mn1.space_before = Pt(36)

# -------------------------------------------------------------
# SLIDE 2: CONTENTS (Exact position: x=1.15", y=1.12", size=18pt)
# -------------------------------------------------------------
s2 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s2, 'CONTENTS')

tb_c = s2.shapes.add_textbox(Inches(1.15), Inches(1.12), Inches(10.5), Inches(6.0))
tf_c = tb_c.text_frame
tf_c.word_wrap = True

contents = [
    '1. Title Slide',
    '2. Content',
    '3. Introduction',
    '4. Technology Used',
    '5. Base Paper / Research Reference',
    '6. Requirement',
    '7. Design (System Architecture)',
    '8. Telemetry & Conjunction Data Schema',
    '9. Class & Component Architecture',
    '10. Development',
    '11. Project Outcome',
    '12. Conclusion',
    '13. Thank you!'
]
for idx, c in enumerate(contents):
    p = tf_c.paragraphs[0] if idx == 0 else tf_c.add_paragraph()
    p.text = c
    p.font.name = 'Arial'
    p.font.size = Pt(18)
    p.font.color.rgb = COLOR_BLACK
    if idx > 0:
        p.space_before = Pt(4)

# -------------------------------------------------------------
# SLIDE 3: INTRODUCTION (Exact position: x=1.07", y=1.20")
# -------------------------------------------------------------
s3 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s3, 'INTRODUCTION')

tb_intro = s3.shapes.add_textbox(Inches(1.07), Inches(1.20), Inches(11.0), Inches(5.8))
tf_in = tb_intro.text_frame
tf_in.word_wrap = True

render_section(tf_in, 'Problem', [
    'Over 36,000 tracked debris objects and constellations create thousands of close conjunctions weekly in LEO.',
    'Missing a real collision event (False Negative) leads to permanent satellite loss and cascading Kessler syndrome.',
    'Conventional manual screening of Conjunction Data Messages (CDMs) is slow, fatigue-prone, and unscalable.'
], first=True, head_size=20, bullet_size=15.5, space_head=0, space_item=3)

render_section(tf_in, 'Objective', [
    'Develop Prahari: an autonomous, high-recall spacecraft conjunction triage and decision-support system.',
    'Predict collision probability log10(Pc) with high precision and actionable early warning across multi-pass CDMs.',
    'Eliminate false alarm floods and provide interactive 3D WebGL orbit visualization for flight safety operators.'
], first=False, head_size=20, bullet_size=15.5, space_head=10, space_item=3)

render_section(tf_in, 'Solution Impact', [
    'Achieves 92.70% High-Risk Recall (165/178 dangerous collisions detected at the official -6.0 alert threshold).',
    'Reduces false alarms by ~75% (79.71% precision, only 42 false alarms vs 85% in published literature).',
    'Provides 3 to 5 days of actionable lead time for Collision Avoidance Maneuver (CAM) thruster burn execution.',
    'Enables automated multi-tier triage and real-time 3D orbital encounter screening.'
], first=False, head_size=20, bullet_size=15.5, space_head=10, space_item=3)

# -------------------------------------------------------------
# SLIDE 4: TECHNOLOGY USED (Exact position: x=1.15", y=1.20")
# -------------------------------------------------------------
s4 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s4, 'TECHNOLOGY USED')

tb_tech = s4.shapes.add_textbox(Inches(1.15), Inches(1.20), Inches(11.0), Inches(5.8))
tf_tech = tb_tech.text_frame
tf_tech.word_wrap = True

render_section(tf_tech, 'Technology Used:', [
    'Frontend: React 18, TypeScript, Vite & Tailwind CSS',
    'Backend: FastAPI (Python 3.11) & Uvicorn ASGI Server',
    'Machine Learning: Extreme Gradient Boosting (XGBoost) Regressor',
    '3D Graphics: Three.js, WebGL & Canvas (Orbital Encounter Scene)',
    'Charts & State: Recharts & TanStack React Query',
    'Optimization: Asymmetric Sample Weighting (w=30.0 for High-Risk Recall Focus)',
    'Data Imputation: Pre-Computed Median Matrix (100% Zero-Crash Reliability)',
    'Data: ESA Kelvins Spacecraft Collision Dataset & Space Surveillance Network CDMs',
    'Tools: Git, GitHub, REST APIs & Windows START.bat Launcher'
], first=True, head_size=20, bullet_size=15.5, space_head=0, space_item=3)

p_bph = tf_tech.add_paragraph()
p_bph.text = 'Base Paper :'
p_bph.font.name = 'Arial'
p_bph.font.size = Pt(20)
p_bph.font.bold = True
p_bph.font.color.rgb = COLOR_BLACK
p_bph.space_before = Pt(12)

p_bp1 = tf_tech.add_paragraph()
p_bp1.text = '“Machine Learning Approaches for Spacecraft Collision Avoidance Challenge”'
p_bp1.font.name = 'Arial'
p_bp1.font.size = Pt(16.5)
p_bp1.font.color.rgb = COLOR_BLACK
p_bp1.space_before = Pt(3)

p_bp2 = tf_tech.add_paragraph()
p_bp2.text = 'Authors: Kalyanaraman et al. (EPJ Web of Conferences, 2026) & Uriot et al. (ESA Advanced Concepts Team)'
p_bp2.font.name = 'Arial'
p_bp2.font.size = Pt(15.5)
p_bp2.font.color.rgb = COLOR_BLACK
p_bp2.space_before = Pt(3)

# -------------------------------------------------------------
# SLIDE 5: BASE PAPER / RESEARCH REFERENCE (Exact position: top=1.20", columns top=3.15")
# -------------------------------------------------------------
s5 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s5, 'BASE PAPER / RESEARCH REFERENCE')

tb_bp_top = s5.shapes.add_textbox(Inches(1.18), Inches(1.20), Inches(11.0), Inches(1.7))
tf_bpt = tb_bp_top.text_frame
tf_bpt.word_wrap = True

p_pt = tf_bpt.paragraphs[0]
p_pt.text = 'Paper Title:'
p_pt.font.name = 'Arial'
p_pt.font.size = Pt(20)
p_pt.font.bold = True
p_pt.font.color.rgb = COLOR_BLACK

p_p1 = tf_bpt.add_paragraph()
p_p1.text = '“Machine Learning Approaches for Spacecraft Collision Avoidance Challenge”'
p_p1.font.name = 'Arial'
p_p1.font.size = Pt(17.5)
p_p1.font.color.rgb = COLOR_BLACK
p_p1.space_before = Pt(2)

p_p2 = tf_bpt.add_paragraph()
p_p2.text = 'Focus: Spacecraft Collision Risk Prediction, CDM Assessment & Orbit Uncertainty'
p_p2.font.name = 'Arial'
p_p2.font.size = Pt(16.5)
p_p2.font.color.rgb = COLOR_BLACK
p_p2.space_before = Pt(2)

p_p3 = tf_bpt.add_paragraph()
p_p3.text = 'Authors: Kalyanaraman et al. (2026) & Uriot et al. (ESA Advanced Concepts Team)'
p_p3.font.name = 'Arial'
p_p3.font.size = Pt(16.5)
p_p3.font.color.rgb = COLOR_BLACK
p_p3.space_before = Pt(2)

# Two Open Columns
tb_ex = s5.shapes.add_textbox(Inches(1.18), Inches(3.15), Inches(5.5), Inches(4.0))
tf_ex = tb_ex.text_frame
tf_ex.word_wrap = True

p_ex_h = tf_ex.paragraphs[0]
p_ex_h.text = 'Existing approach :'
p_ex_h.font.name = 'Arial'
p_ex_h.font.size = Pt(20)
p_ex_h.font.bold = True
p_ex_h.font.color.rgb = COLOR_BLACK

ex_points = [
    'Static single-snapshot CDM evaluation',
    'Evaluated strictly at TCA (t = 0 hours)',
    'Standard binary classification with SMOTE',
    'Assumes equal cost between False Positives and False Negatives',
    'Black-box feature engineering with opaque synthetic indices'
]
for pt in ex_points:
    p = tf_ex.add_paragraph()
    p.text = '•  ' + pt
    p.font.name = 'Arial'
    p.font.size = Pt(15.5)
    p.font.color.rgb = COLOR_BLACK
    p.space_before = Pt(5)

tb_rg = s5.shapes.add_textbox(Inches(7.15), Inches(3.15), Inches(5.5), Inches(4.0))
tf_rg = tb_rg.text_frame
tf_rg.word_wrap = True

p_rg_h = tf_rg.paragraphs[0]
p_rg_h.text = 'Research Gap :'
p_rg_h.font.name = 'Arial'
p_rg_h.font.size = Pt(20)
p_rg_h.font.bold = True
p_rg_h.font.color.rgb = COLOR_BLACK

rg_points = [
    'Temporal Blindness: 0-hour lead time is useless for thruster burn prep -> Solved with 3-5 days lead time.',
    'False Alarm Flooding: 85% false alarm rate (15% precision) exhausts fuel -> Solved with 79.71% precision.',
    'Class Imbalance: 178 collisions vs 1,989 safe events -> Solved via 30x Asymmetric Weighting.',
    'High Recall Focus: Prioritizing recall (92.70%) to ensure spacecraft survival.',
    'Native Telemetry: Ingests 100 raw CDM features with 0 synthetic bloat.'
]
for pt in rg_points:
    p = tf_rg.add_paragraph()
    p.text = '•  ' + pt
    p.font.name = 'Arial'
    p.font.size = Pt(15.5)
    p.font.color.rgb = COLOR_BLACK
    p.space_before = Pt(5)

# -------------------------------------------------------------
# SLIDE 6: REQUIREMENT (Exact position: x=1.07", y=1.20")
# -------------------------------------------------------------
s6 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s6, 'REQUIREMENT')

tb_req = s6.shapes.add_textbox(Inches(1.07), Inches(1.20), Inches(11.0), Inches(5.8))
tf_req = tb_req.text_frame
tf_req.word_wrap = True

p_rh = tf_req.paragraphs[0]
p_rh.text = 'Product:'
p_rh.font.name = 'Arial'
p_rh.font.size = Pt(20)
p_rh.font.bold = True
p_rh.font.color.rgb = COLOR_BLACK

p_rp = tf_req.add_paragraph()
p_rp.text = 'Prahari: Autonomous Spacecraft Conjunction Triage & Collision Avoidance Decision-Support System'
p_rp.font.name = 'Arial'
p_rp.font.size = Pt(16.5)
p_rp.font.color.rgb = COLOR_BLACK
p_rp.space_before = Pt(2)

render_section(tf_req, 'Features:', [
    'Automated Multi-Pass CDM Ingestion & Telemetry Preprocessing',
    'High-Recall XGBoost Conjunction Risk Regression (log10 Pc prediction)',
    'Asymmetric Risk Thresholding calibrated at the official -6.0 alert boundary',
    'Real-time Multi-Tier Action Protocol (CRITICAL, HIGH, ELEVATED, LOW bands)',
    '100 Native Telemetry Directory (Kinematics, Covariance, Solar Weather, Tracking)',
    'Interactive 3D WebGL Orbit & Covariance Ellipsoid Encounter Visualizer',
    'Temporal Timeline Playback Scrubber for multi-day radar update evolution',
    'Model Lab with live threshold calibration slider and SHAP feature importance charts',
    'Drag-and-Drop Live Telemetry Prediction Sandbox for operational passes',
    'Automated Collision Avoidance Maneuver (CAM) alert and report generation'
], first=False, head_size=20, bullet_size=15.0, space_head=10, space_item=3.0)

# -------------------------------------------------------------
# SLIDE 7: DESIGN (SYSTEM ARCHITECTURE)
# -------------------------------------------------------------
s7 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s7, 'DESIGN')

tb_d1 = s7.shapes.add_textbox(Inches(1.07), Inches(0.98), Inches(11.0), Inches(0.5))
tf_d1 = tb_d1.text_frame
p_d1 = tf_d1.paragraphs[0]
p_d1.text = 'Prahari – System Design & Pipeline Architecture'
p_d1.font.name = 'Arial'
p_d1.font.size = Pt(20)
p_d1.font.bold = True
p_d1.font.color.rgb = COLOR_BLACK
p_d1.alignment = PP_ALIGN.CENTER

diag1_file = os.path.join(DOCS_DIR, 'system_design_diagram.png')
if os.path.exists(diag1_file):
    s7.shapes.add_picture(diag1_file, Inches(1.4), Inches(1.6), Inches(10.53), Inches(5.4))

# -------------------------------------------------------------
# SLIDE 8: DESIGN (DATABASE & TELEMETRY SCHEMA)
# -------------------------------------------------------------
s8 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s8, 'DESIGN')

tb_d2 = s8.shapes.add_textbox(Inches(1.07), Inches(0.98), Inches(11.0), Inches(0.5))
tf_d2 = tb_d2.text_frame
p_d2 = tf_d2.paragraphs[0]
p_d2.text = 'Prahari – Telemetry & Conjunction Data Schema'
p_d2.font.name = 'Arial'
p_d2.font.size = Pt(20)
p_d2.font.bold = True
p_d2.font.color.rgb = COLOR_BLACK
p_d2.alignment = PP_ALIGN.CENTER

diag2_file = os.path.join(DOCS_DIR, 'database_diagram.png')
if os.path.exists(diag2_file):
    s8.shapes.add_picture(diag2_file, Inches(1.4), Inches(1.6), Inches(10.53), Inches(5.4))

# -------------------------------------------------------------
# SLIDE 9: DESIGN (CLASS & COMPONENT ARCHITECTURE)
# -------------------------------------------------------------
s9 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s9, 'DESIGN')

tb_d3 = s9.shapes.add_textbox(Inches(1.07), Inches(0.98), Inches(11.0), Inches(0.5))
tf_d3 = tb_d3.text_frame
p_d3 = tf_d3.paragraphs[0]
p_d3.text = 'Prahari – Component & Class Architecture'
p_d3.font.name = 'Arial'
p_d3.font.size = Pt(20)
p_d3.font.bold = True
p_d3.font.color.rgb = COLOR_BLACK
p_d3.alignment = PP_ALIGN.CENTER

diag3_file = os.path.join(DOCS_DIR, 'class_diagram.png')
if os.path.exists(diag3_file):
    s9.shapes.add_picture(diag3_file, Inches(1.4), Inches(1.6), Inches(10.53), Inches(5.4))

# -------------------------------------------------------------
# SLIDE 10: DEVELOPMENT (Exact position: x=1.15", y=1.20")
# -------------------------------------------------------------
s10 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s10, 'DEVELOPMENT')

tb_dev = s10.shapes.add_textbox(Inches(1.15), Inches(1.20), Inches(11.0), Inches(5.8))
tf_dev = tb_dev.text_frame
tf_dev.word_wrap = True

render_section(tf_dev, 'Development Steps:', [
    'Extracted and audited 100 native telemetry parameters across 162,634 observation rows.',
    'Engineered asymmetric sample weighting protocol (w = 30.0 for high risk) to prioritize recall.',
    'Trained regularized XGBoost Regressor with histogram binning and pre-computed median imputation.',
    'Evaluated performance on 2,167 completely held-out test events at the standard -6.0 threshold.',
    'Built asynchronous FastAPI backend serving sub-50ms real-time inference and metadata endpoints.',
    'Developed React 18 + Vite frontend with Three.js 3D WebGL orbital encounter scene and timeline scrubber.',
    'Created interactive Model Lab with dynamic threshold slider, 100-feature directory, and SHAP charts.',
    'Authored publication-ready academic research paper (.docx) and 16:9 presentation deck (.pptx).'
], first=True, head_size=20, bullet_size=15.5, space_head=0, space_item=3.0)

render_section(tf_dev, 'Tools Used:', [
    'React, TypeScript & Vite',
    'Tailwind CSS & Lucide Icons',
    'Python 3.11, FastAPI & Uvicorn',
    'XGBoost, Scikit-learn, Pandas & Joblib',
    'Three.js & WebGL Canvas Renderer',
    'Git & GitHub'
], first=False, head_size=20, bullet_size=15.5, space_head=10, space_item=3.0)

# -------------------------------------------------------------
# SLIDE 11: PROJECT OUTCOME (Exact position: x=1.15", y=1.20")
# -------------------------------------------------------------
s11 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s11, 'PROJECT OUTCOME')

tb_out = s11.shapes.add_textbox(Inches(1.15), Inches(1.20), Inches(11.0), Inches(5.8))
tf_out = tb_out.text_frame
tf_out.word_wrap = True

render_section(tf_out, 'Achievements', [
    'High-Risk Recall: 92.70% (165 of 178 dangerous collisions detected ahead of time).',
    'Safety F2 Score: 89.77% (Recall weighted 2x to prioritize mission survival).',
    'Precision: 79.71% (Only 42 false alarms across 1,989 safe events, eliminating the 85% false alarm flood).',
    'Overall Accuracy: 97.46% (2,112 out of 2,167 events correctly classified).',
    '100% Native Telemetry: Ingests raw CDMs directly without artificial synthetic features.',
    'Actionable Lead Time: 3 to 5 days pre-TCA for Collision Avoidance Maneuver (CAM) execution.'
], first=True, head_size=20, bullet_size=15.0, space_head=0, space_item=2.5)

render_section(tf_out, 'Current Challenges', [
    'Sparse ground radar coverage causes intermittent tracking gaps in higher orbital inclinations.',
    'Solar storm geomagnetic surges temporarily inflate atmospheric drag covariance errors.'
], first=False, head_size=20, bullet_size=15.0, space_head=9, space_item=2.5)

render_section(tf_out, 'Future Improvements', [
    'Autonomous CAM Delta-V burn calculation using Reinforcement Learning.',
    'Live Space-Track / ESA API integration for continuous automatic ingest.',
    'Constellation multi-satellite screening to avoid secondary induced conjunctions.'
], first=False, head_size=20, bullet_size=15.0, space_head=9, space_item=2.5)

p_paph = tf_out.add_paragraph()
p_paph.text = 'Paper:'
p_paph.font.name = 'Arial'
p_paph.font.size = Pt(20)
p_paph.font.bold = True
p_paph.font.color.rgb = COLOR_BLACK
p_paph.space_before = Pt(9)

p_pap1 = tf_out.add_paragraph()
p_pap1.text = 'Status: Research paper is complete and formatted for journal submission at docs/Prahari_Research_Paper.docx.'
p_pap1.font.name = 'Arial'
p_pap1.font.size = Pt(15.0)
p_pap1.font.color.rgb = COLOR_BLACK
p_pap1.space_before = Pt(2)

# -------------------------------------------------------------
# SLIDE 12: CONCLUSION (Exact position: x=1.15", y=1.20")
# -------------------------------------------------------------
s12 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(s12, 'CONCLUSION')

tb_con = s12.shapes.add_textbox(Inches(1.15), Inches(1.20), Inches(11.0), Inches(5.8))
tf_con = tb_con.text_frame
tf_con.word_wrap = True

render_section(tf_con, 'Key Takeaways', [
    'Prahari transforms multi-pass raw CDM telemetry into high-recall, actionable collision warnings.',
    'Asymmetric loss optimization successfully solves the class imbalance and catches 92.70% of collisions.',
    'Eliminates the temporal blindness (TCA=0h) and 85% false alarm rate of published literature.',
    'Interactive 3D Three.js mission control bridges machine learning with operational flight dynamics.'
], first=True, head_size=20, bullet_size=15.5, space_head=0, space_item=3.0)

render_section(tf_con, 'New Skills Learned', [
    'Space Situational Awareness (SSA) & Astrodynamic Conjunction Assessment',
    'Asymmetric Loss Formulation for Extreme Class Imbalance in Space Systems',
    'High-Dimensional Telemetry Engineering on 100 Native Parameters',
    'Full-Stack Modern Web Development (React 18, TypeScript, Vite, Tailwind CSS)',
    '3D Orbital Graphics & Covariance Ellipsoid Rendering (Three.js & WebGL)',
    'High-Performance Asynchronous Microservices (FastAPI & Uvicorn)',
    'Academic Research Paper Authoring & Rigorous Empirical Benchmarking'
], first=False, head_size=20, bullet_size=15.5, space_head=10, space_item=3.0)

# -------------------------------------------------------------
# SLIDE 13: THANK YOU!
# -------------------------------------------------------------
s13 = prs.slides.add_slide(prs.slide_layouts[6])
apply_slide_background(s13, is_title=False)

tb_ty = s13.shapes.add_textbox(Inches(0), Inches(3.2), Inches(13.333), Inches(1.5))
tf_ty = tb_ty.text_frame
p_ty = tf_ty.paragraphs[0]
p_ty.text = 'Thank you!'
p_ty.font.name = 'Arial'
p_ty.font.size = Pt(48)
p_ty.font.bold = True
p_ty.font.color.rgb = COLOR_BLACK
p_ty.alignment = PP_ALIGN.CENTER

# -------------------------------------------------------------
# SAVE PRESENTATION & EXPORT PREVIEWS
# -------------------------------------------------------------
output_path = os.path.join(DOCS_DIR, 'Prahari_Presentation.pptx')
prs.save(output_path)
print(f'Successfully generated presentation at: {output_path}')

# Export previews
try:
    ppt_app = win32com.client.Dispatch('PowerPoint.Application')
    pres = ppt_app.Presentations.Open(output_path, WithWindow=False)
    for i, slide in enumerate(pres.Slides):
        img_path = os.path.join(PREVIEWS_DIR, f'slide_{i+1}.png')
        slide.Export(img_path, 'PNG', 1920, 1080)
    pres.Close()
    ppt_app.Quit()
    print('All 13 slides exported to PNG successfully!')
except Exception as e:
    print('PowerPoint export error:', e)
